import io
import json
import os
import sqlite3
import uuid
import zipfile
import time
from datetime import datetime, timezone
from pathlib import Path
from contextlib import contextmanager
from dotenv import load_dotenv
from fastapi import FastAPI, HTTPException, UploadFile, File, Form, Request
from fastapi.responses import FileResponse, StreamingResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from docx import Document
from docx.shared import Pt, Inches
from .models import NewTask, Extraction, Review
from .extract import parse, extract
from .samples import sample, LABELS
from .rules import audit, KINDS, number
from .word import style_document
from .settings import load_settings
from .observability import emit, request_id, task_calls
from .auth import install_auth, session_user, same_origin
from .intelligence import analyze

ROOT=Path(__file__).resolve().parents[1]
load_dotenv(ROOT/'.env')
DATA=Path(os.getenv('DATA_DIR',str(ROOT/'data')));DATA.mkdir(parents=True,exist_ok=True)
UPLOAD=DATA/'uploads';UPLOAD.mkdir(exist_ok=True)
SAMPLES=DATA/'samples';SAMPLES.mkdir(exist_ok=True)
app=FastAPI(title='采核助手', version='1.0.0')

@app.middleware('http')
async def request_logging(request:Request,call_next):
    if not request.url.path.startswith('/api/'):
        return await call_next(request)
    rid=uuid.uuid4().hex;token=request_id.set(rid);start=time.perf_counter();status=500
    try:
        public={'/api/auth/login','/api/auth/logout','/api/auth/session','/api/health'}
        if request.url.path not in public and not session_user(request,db):
            response=JSONResponse({'detail':'请先登录后再访问'},401)
        elif request.method not in ['GET','HEAD','OPTIONS'] and not same_origin(request):
            response=JSONResponse({'detail':'请求来源无效'},403)
        else:
            response=await call_next(request)
        response.headers['Cache-Control']='no-store'
        status=response.status_code
        response.headers['X-Request-ID']=rid
        return response
    finally:
        # ASGI route templates exclude user-provided path text and query values.
        route=request.scope.get('route');path=getattr(route,'path','/api/unmatched')
        emit('http.completed',method=request.method,path=path,status=status,duration_ms=round((time.perf_counter()-start)*1000,1))
        request_id.reset(token)

@contextmanager
def db():
    con=sqlite3.connect(DATA/'app.db');con.row_factory=sqlite3.Row
    try: yield con;con.commit()
    finally: con.close()

with db() as con:
    con.execute('CREATE TABLE IF NOT EXISTS tasks (id TEXT PRIMARY KEY, body TEXT NOT NULL)')

install_auth(app,db)

def now(): return datetime.now(timezone.utc).isoformat()
def save(t,create_new=False):
    t['updated_at']=now()
    with db() as con:
        if create_new:con.execute('INSERT INTO tasks VALUES (?,?)',(t['id'],json.dumps(t,ensure_ascii=False)))
        elif con.execute('UPDATE tasks SET body=? WHERE id=?',(json.dumps(t,ensure_ascii=False),t['id'])).rowcount!=1:
            raise HTTPException(404,'任务已删除，未保存过期结果')
def get(tid):
    with db() as con: row=con.execute('SELECT body FROM tasks WHERE id=?',(tid,)).fetchone()
    if not row: raise HTTPException(404,'任务不存在')
    return json.loads(row['body'])
def invalidate(t,fields_changed=False):
    if t.get('result'): t.setdefault('history',[]).append(t['result'])
    t['result']=None
    t.pop('review_confirmation',None)
    t.pop('ai_analysis',None)
    if fields_changed: t.pop('field_confirmation',None)

def require_submission(t):
    if not t.get('materials_submission'):
        raise HTTPException(409,'请先上传三类材料并提交，再进行后续操作')

@app.get('/api/config')
def config():
    try:
        settings=load_settings()
    except ValueError as e:
        return dict(mode='配置错误',max_size_mb=10,error=str(e))
    return dict(mode=('DeepSeek 模型模式' if settings.provider=='deepseek' else '模型模式') if settings.configured else '演示模式',
                provider=settings.provider,model=settings.name,configured=settings.configured,max_size_mb=10)

@app.get('/api/tasks')
def tasks():
    with db() as con: rows=con.execute('SELECT body FROM tasks ORDER BY rowid DESC').fetchall()
    return [json.loads(r['body']) for r in rows]

@app.post('/api/tasks')
def create(body:NewTask):
    t=dict(id=uuid.uuid4().hex,name=body.name,created_at=now(),documents=[],edits=[],result=None,demo=False)
    save(t,create_new=True);return t

@app.get('/api/tasks/{tid}')
def task(tid:str): return get(tid)

@app.delete('/api/tasks/{tid}')
def delete_task(tid:str):
    # Resolve owned files before touching them; never delete shared sample assets.
    with db() as con:
        con.execute('BEGIN IMMEDIATE')
        row=con.execute('SELECT body FROM tasks WHERE id=?',(tid,)).fetchone()
        if not row:raise HTTPException(404,'任务不存在或已删除')
        t=json.loads(row['body'])
        others=[json.loads(r['body']) for r in con.execute('SELECT body FROM tasks WHERE id<>?',(tid,)).fetchall()]
        references={str(Path(d['path']).resolve()) for other in others for d in other['documents']}
        owned=set()
        for d in t['documents']:
            path=Path(d['path']).resolve()
            if path.is_relative_to(UPLOAD.resolve()) and str(path) not in references:owned.add(path)
        con.execute('DELETE FROM tasks WHERE id=?',(tid,))
    remaining=0
    for path in owned:
        try:path.unlink(missing_ok=True)
        except OSError:remaining+=1
    emit('task.deleted',task_id=tid,deleted_uploads=len(owned)-remaining,cleanup_pending=remaining)
    return dict(deleted=True,id=tid,deleted_uploads=len(owned)-remaining,cleanup_pending=remaining)

@app.get('/api/tasks/{tid}/calls')
def calls(tid:str):
    get(tid)
    return task_calls(tid)

@app.post('/api/tasks/{tid}/submit-materials')
def submit_materials(tid:str):
    t=get(tid)
    if len(t['documents'])!=3 or set(d['kind'] for d in t['documents'])!=set(KINDS):
        raise HTTPException(409,'请先上传采购需求、供应商报价单和采购合同')
    if any(not d['blocks'] for d in t['documents']):
        raise HTTPException(409,'存在无法解析的文件，请替换为可提取文本的材料后提交')
    if not t.get('materials_submission'): invalidate(t,fields_changed=True)
    t['materials_submission']=now();save(t);return t

@app.post('/api/tasks/{tid}/confirm-fields')
def confirm_fields(tid:str):
    t=get(tid)
    require_submission(t)
    if len(t['documents'])!=3 or any(d['status']!='已抽取' for d in t['documents']):
        raise HTTPException(400,'请先上传三类材料并完成抽取或人工录入')
    t['field_confirmation']=now();t.pop('review_confirmation',None);save(t);return t

@app.post('/api/tasks/{tid}/confirm-review')
def confirm_review(tid:str):
    t=get(tid)
    require_submission(t)
    if not t.get('field_confirmation') or not t['result']:
        raise HTTPException(400,'请先确认抽取字段并执行核对')
    if any(i['status']=='待处理' for i in t['result']['issues']):
        raise HTTPException(400,'请逐项处理问题，再完成人工复核')
    if any(i['status']=='待处理' for i in (t.get('ai_analysis') or {}).get('findings',[])):
        raise HTTPException(400,'AI 辅助分析仍有待处理建议，请先人工确认')
    t['review_confirmation']=now();save(t);return t

@app.post('/api/tasks/{tid}/documents')
async def upload(tid:str,kind:str=Form(...),file:UploadFile=File(...)):
    t=get(tid)
    if kind not in KINDS: raise HTTPException(400,'请选择正确的文档类型')
    if any(d['kind']==kind for d in t['documents']): raise HTTPException(400,'该类型文档已存在，请先删除后重新上传')
    suffix=Path(file.filename or '').suffix.lower()
    if suffix not in ['.pdf','.docx']: raise HTTPException(400,'仅支持可提取文本的 PDF 和 DOCX 文件')
    content=await file.read(10*1024*1024+1)
    if len(content)>10*1024*1024: raise HTTPException(400,'文件超过 10 MB 上限')
    did=uuid.uuid4().hex;path=UPLOAD/(did+suffix);path.write_bytes(content)
    d=dict(id=did,name=Path(file.filename).name,kind=kind,path=str(path),status='未抽取',blocks=[],data=None,error=None,demo=False)
    try: d['blocks']=parse(path)
    except Exception as e:
        d.update(status='抽取失败',error=str(e) if isinstance(e,ValueError) else '文档解析失败：文件损坏、加密或格式不正确')
    t['documents'].append(d);t.pop('materials_submission',None);invalidate(t,fields_changed=True);save(t);return t

@app.delete('/api/tasks/{tid}/documents/{did}')
def delete_doc(tid:str,did:str):
    t=get(tid);t['documents']=[d for d in t['documents'] if d['id']!=did];t.pop('materials_submission',None);invalidate(t,fields_changed=True);save(t);return t

@app.post('/api/tasks/{tid}/extract')
async def do_extract(tid:str):
    t=get(tid);require_submission(t);invalidate(t,fields_changed=True)
    for d in t['documents']:
        if d['demo'] or d['status']=='已抽取' or not d['blocks']:
            emit('model.skipped',task_id=tid,document_id=d['id'],reason='demo_data' if d['demo'] else 'already_extracted' if d['status']=='已抽取' else 'parse_failed')
            continue
        try:
            d['data']=await extract(d['blocks'],task_id=tid,document_id=d['id'])
            d.update(status='已抽取',error=None,extraction_source='model',extracted_at=now())
        except ValueError as e: d.update(status='抽取失败',error=str(e))
    save(t);return t

@app.put('/api/tasks/{tid}/documents/{did}')
def edit(tid:str,did:str,body:Extraction):
    t=get(tid);require_submission(t);d=next((d for d in t['documents'] if d['id']==did),None)
    if not d: raise HTTPException(404,'文档不存在')
    if not d['blocks']: raise HTTPException(400,'无法解析的文件不能手工覆盖，请上传可提取文本的文件')
    new=body.model_dump()
    # Human assertions remain distinct from source quotations; validate any cited text.
    from .extract import validate_evidence
    try:
        def validate_manual(obj):
            if isinstance(obj,dict):
                if 'value' in obj:
                    if obj['quote']: validate_evidence(obj,d['blocks'])
                    elif obj['value'] is not None: obj['location']='人工录入 无原文引用'
                else:
                    for v in obj.values(): validate_manual(v)
            elif isinstance(obj,list):
                for v in obj: validate_manual(v)
        validate_manual(new)
    except ValueError as e: raise HTTPException(400,str(e))
    t['edits'].append(dict(document_id=did,time=now(),before=d['data'],after=new))
    if d['data'] is None:d['extraction_source']='manual'
    d.update(data=new,status='已抽取',error=None,manually_modified=True);invalidate(t,fields_changed=True);save(t);return t

@app.post('/api/tasks/{tid}/audit')
def run_audit(tid:str,tolerance:str='0.01'):
    t=get(tid)
    require_submission(t)
    if not t.get('field_confirmation') or len(t['documents'])!=3 or any(d['status']!='已抽取' for d in t['documents']):
        raise HTTPException(409,'请先完成三类材料的抽取并人工确认字段，再执行核对')
    try: result=audit(t['documents'],tolerance)
    except ValueError as e: raise HTTPException(400,str(e))
    invalidate(t);t['result']=result;save(t);return t

@app.post('/api/tasks/{tid}/ai-analysis')
async def ai_analysis(tid:str):
    t=get(tid);require_submission(t)
    if not t.get('field_confirmation') or not t['result']:
        raise HTTPException(409,'请先确认字段并执行代码规则核对，再进行 AI 语义复核')
    version=t.get('updated_at')
    try:result=await analyze(t)
    except ValueError as exc:result={'status':'失败','error':str(exc),'findings':[],'source':'model'}
    current=get(tid)
    if current.get('updated_at')!=version:raise HTTPException(409,'分析期间任务已变更，本次 AI 结果未采用，请重新分析')
    if t.get('ai_analysis'):t.setdefault('ai_analysis_history',[]).append(t['ai_analysis'])
    t['ai_analysis']={**result,'created_at':now()};t.pop('review_confirmation',None)
    save(t);return t

@app.put('/api/tasks/{tid}/ai-findings/{iid}')
def review_ai(tid:str,iid:str,body:Review):
    t=get(tid);analysis=t.get('ai_analysis') or {}
    finding=next((f for f in analysis.get('findings',[]) if f['id']==iid),None)
    if analysis.get('status')!='已完成' or not finding:raise HTTPException(404,'AI 建议不存在或已失效')
    finding.update(**body.model_dump());t.pop('review_confirmation',None)
    t.setdefault('review_log',[]).append(dict(issue=iid,source='ai',time=now(),**body.model_dump()))
    save(t);return t

@app.put('/api/tasks/{tid}/issues/{iid}')
def review(tid:str,iid:str,body:Review):
    t=get(tid)
    if not t['result']: raise HTTPException(400,'请先执行审核')
    issue=next((i for i in t['result']['issues'] if i['id']==iid),None)
    if not issue: raise HTTPException(404,'问题不存在')
    issue.update(**body.model_dump());t.pop('review_confirmation',None);t.setdefault('review_log',[]).append(dict(issue=iid,time=now(),**body.model_dump()));save(t);return t

@app.get('/api/tasks/{tid}/documents/{did}/file')
def source(tid:str,did:str):
    d=next((d for d in get(tid)['documents'] if d['id']==did),None)
    if not d: raise HTTPException(404,'文档不存在')
    return FileResponse(d['path'],filename=d['name'])

@app.post('/api/samples/{group}')
def load_sample(group:str):
    if group not in ['A','B']: raise HTTPException(404,'样例不存在')
    t=create(NewTask(name=f'{group}组 · 研发办公显示器采购'));t['demo']=True
    for kind in KINDS:
        path,data=sample(group,kind,SAMPLES)
        t['documents'].append(dict(id=uuid.uuid4().hex,name=path.name,kind=kind,path=str(path),status='已抽取',blocks=parse(path),data=data,error=None,demo=True,extraction_source='demo'))
    t['materials_submission']=now()
    save(t);return t

@app.get('/api/samples/{group}/download')
def download_samples(group:str):
    if group not in ['A','B']: raise HTTPException(404,'样例不存在')
    out=io.BytesIO()
    with zipfile.ZipFile(out,'w') as z:
        for kind in KINDS:
            path,_=sample(group,kind,SAMPLES);z.write(path,path.name)
        expected='A组：无预设异常。\n' if group=='A' else 'B组人工标注：\nITEM-001 型号：需求与报价 MX27；合同 MX24。\nITEM-002 数量：需求与报价 10台；合同12台。\nMATH-001：合同12×600=7200，但分项金额6500。\nDATE-001：需求和报价最晚2026-11-30；合同2026-12-15。\nFIELD-001：合同缺少税费口径。\nBASIS-001：合同税费未知，不直接比较报价合同金额。\n'
        z.writestr('预期问题清单.txt',expected+'预设抽取仅用于演示，不代表真实模型准确率。')
    out.seek(0);return StreamingResponse(out,media_type='application/zip',headers={'Content-Disposition':f'attachment; filename=sample-{group}.zip'})

@app.get('/api/tasks/{tid}/report')
def report(tid:str):
    t=get(tid);r=t['result']
    require_submission(t)
    if not t.get('field_confirmation'): raise HTTPException(409,'请先确认抽取字段，再生成审核报告')
    if not r: raise HTTPException(409,'尚未执行核对，审核报告暂未生成')
    doc=style_document(Document())
    sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    style=doc.styles['Normal'];style.font.name='Microsoft YaHei';style.font.size=Pt(11)
    style.paragraph_format.space_after=Pt(4);style.paragraph_format.line_spacing=1.1
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    style.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    doc.add_heading('采购材料交叉审核报告',0)
    doc.add_paragraph(f"本报告针对 {t['name']}，核对采购需求、报价单和合同中的业务信息；不构成法律合规审查。自动结果：{r['status'] if r else '未检查'}。人工处理状态单独列示。")
    if r:
        counts={s:sum(i['status']==s for i in r['issues']) for s in ['待处理','已确认','误报']}
        doc.add_paragraph(f"自动发现 {len(r['issues'])} 项问题。人工处理：待处理 {counts['待处理']} 项，已确认 {counts['已确认']} 项，误报 {counts['误报']} 项。人工判断不改变原自动检测结果。")
    doc.add_heading('任务信息',1)
    doc.add_paragraph(f"抽取字段人工确认：{t.get('field_confirmation') or '未确认'}\n整体人工复核：{t.get('review_confirmation') or '未完成'}。确认不代表所有异常已消除，仍需查看自动结果及未完成项。")
    doc.add_paragraph(f"任务编号：{tid}\n创建时间 UTC：{t['created_at']}\n数据来源：{'内置演示数据，非真实模型准确率' if t['demo'] else '用户上传'}\n金额允许误差：{r['tolerance'] if r else '未检查'}")
    doc.add_heading('文档清单与审核范围',1)
    doc.add_paragraph('范围：必要材料与字段、金额运算、商品型号数量单位、供应商及交付要求。仅成功抽取的内容进入字段审核；相对日期和不明口径需人工复核。')
    for d in t['documents']: doc.add_paragraph(f"{d['kind']} · {d['name']} · {d['status']}"+(f"\n{d['error']}" if d.get('error') else ''))
    for k in KINDS:
        if not any(d['kind']==k for d in t['documents']): doc.add_paragraph(f'{k}：未获取，该类文档内容未检查')
    if not t['documents']: doc.add_paragraph('未获取文档')
    doc.add_heading('问题汇总与原文依据',1)
    if not r: doc.add_paragraph('未检查，不得据此得出通过结论。')
    elif not r['issues']: doc.add_paragraph('自动规则未发现问题；审核范围及未完成检查项见下文。')
    else:
        for i in r['issues']:
            doc.add_heading(f"{i['id']} {i['type']} {i['severity']}",2)
            doc.add_paragraph(f"规则：{i['rule']}\n相关字段：{'、'.join(LABELS.get(f,f) for f in i['fields'])}\n解释：{i['explanation']}\n建议：{i['suggestion']}\n人工状态：{i['status']}\n备注：{i['note'] or '未填写'}")
            for e in i['evidence']: doc.add_paragraph(f"文档：{e['document']} ({e['document_id']})\n位置：{e['location'] or '未获取'}\n原文：{e['quote'] or '未获取'}\n字段值：{e['value'] if e['value'] is not None else '未获取'}")
            if not i['evidence']: doc.add_paragraph('原文依据：未获取')
    doc.add_heading('AI 语义复核（独立于代码规则）',1)
    ai=t.get('ai_analysis')
    if not ai:doc.add_paragraph('未检查：尚未执行 AI 语义复核，不得据此得出语义一致结论。')
    elif ai['status']=='失败':doc.add_paragraph('失败，未采用模型结论：'+ai['error'])
    else:
        doc.add_paragraph(f"模型：{ai['model']}\n调用编号：{ai['call_id']}\n摘要：{ai['summary']}")
        def citations(items):
            for e in items:doc.add_paragraph(f"依据：{e['document']} ({e['document_id']}) · {e['location']}\n原文：{e['quote']}")
        citations(ai['summary_evidence'])
        if not ai['findings']:doc.add_paragraph('模型未提出额外建议；不代表审核通过。')
        for f in ai['findings']:
            doc.add_heading(f"{f['id']} {f['title']}",2)
            doc.add_paragraph(f"类型：{f['category']} · 严重程度：{f['severity']}\n解释：{f['explanation']}\n建议：{f['suggestion']}\n关联规则问题：{'、'.join(f['related_issue_ids']) or '无'}\n人工状态：{f['status']}\n备注：{f['note'] or '未填写'}")
            citations(f['evidence'])
        for limitation in ai['limitations']:doc.add_paragraph('分析限制：'+limitation)
    doc.add_heading('未完成检查项',1)
    pending=[c for c in (r['checks'] if r else []) if c['status'] in ['未检查','抽取失败']]
    for c in pending: doc.add_paragraph(f"{c['rule']} {c['status']}：{c['detail']}")
    if not pending: doc.add_paragraph('未执行审核' if not r else '无额外未完成记录；问题中的待复核与商品匹配项仍需人工处理。')
    if t['edits']: doc.add_heading('人工修改记录',1)
    for e in t['edits']:
        doc.add_paragraph(f"{e['time']} 文档 {e['document_id']}")
        def changes(before,after,prefix=''):
            lines=[]
            if isinstance(after,dict) and 'value' in after:
                if before != after: lines.append(f"{prefix}：{(before or {}).get('value') or '未获取'} → {after['value'] or '未获取'}；引用 {after['quote'] or '未获取'} {after['location']}")
            elif isinstance(after,dict):
                for k,v in after.items(): lines+=changes((before or {}).get(k),v,prefix+'.'+k)
            elif isinstance(after,list):
                for idx,v in enumerate(after): lines+=changes(before[idx] if before and idx<len(before) else None,v,f'{prefix}[{idx+1}]')
                if before and len(before)>len(after): lines.append(f'{prefix}：移除 {len(before)-len(after)} 条商品，修改前值已保存在任务修改记录中')
            return lines
        for line in changes(e['before'],e['after']): doc.add_paragraph(line)
    if not t['edits']: doc.add_paragraph('人工修改记录：无')
    out=io.BytesIO();doc.save(out);out.seek(0)
    return StreamingResponse(out,media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',headers={'Content-Disposition':'attachment; filename=audit-report.docx'})

if (ROOT/'frontend/dist').exists(): app.mount('/',StaticFiles(directory=ROOT/'frontend/dist',html=True),name='web')
