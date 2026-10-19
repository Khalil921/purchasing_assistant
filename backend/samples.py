from pathlib import Path
from docx import Document
from .models import Extraction
from .word import style_document

LABELS={'buyer':'采购方','supplier':'供应商','subject':'采购事项','total':'总金额','currency':'币种','tax':'税费口径','delivery':'交付要求','name':'产品名称','model':'型号','quantity':'数量','unit':'单位','price':'单价','amount':'分项金额'}

def sample(group,kind,folder):
    values=dict(buyer='明川科技有限公司',supplier='远山设备有限公司',subject='研发办公显示器采购',total='6000',currency='人民币',tax='含13%增值税',delivery='最晚2026-11-30交付')
    item=dict(name='专业显示器',model='MX27',quantity='10',unit='台',price='600',amount='6000')
    if kind=='采购需求':
        values.update(supplier=None,total=None,currency=None,tax=None)
        item.update(price=None,amount=None)
    if group=='B' and kind=='采购合同':
        values.update(delivery='2026-12-15交付',total='6500',tax=None)
        item.update(model='MX24',quantity='12',amount='6500')
    d=style_document(Document());d.add_heading(f'{group}组 {kind}',0)
    d.add_paragraph('演示数据 用于采核助手业务流程验证')
    out=Extraction().model_dump()
    def write(target,key,value):
        if value is None: return
        text=f'{LABELS[key]}：{value}'
        d.add_paragraph(text)
        target[key]=dict(value=value,quote=text,location=f'DOCX 段落 {len(d.paragraphs)}')
    for k,v in values.items(): write(out,k,v)
    obj=out['items'];obj.append(Extraction(items=[{}]).model_dump()['items'][0])
    for k,v in item.items(): write(obj[0],k,v)
    path=Path(folder)/f'{group}-{kind}.docx';d.save(path)
    return path,out
