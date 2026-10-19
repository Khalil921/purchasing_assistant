import copy
import io
import zipfile
import pytest
from fastapi.testclient import TestClient
from backend.models import Extraction
from backend.rules import audit, matching, date_range, number, KINDS
from backend.samples import sample
from backend.extract import parse, validate_evidence
from backend.main import app

@pytest.fixture
def docs(tmp_path):
    def build(group='A'):
        result=[]
        for n,k in enumerate(KINDS):
            path,data=sample(group,k,tmp_path)
            blocks=parse(path)
            validate_evidence(data,blocks)
            result.append(dict(id=str(n),kind=k,name=path.name,data=data,status='已抽取',blocks=blocks))
        return result
    return build

def rules(ds): return {i['rule'] for i in audit(ds)['issues']}

def test_a_no_preset_anomaly(docs):
    r=audit(docs());assert r['status']=='通过';assert r['issues']==[]

def test_b_all_annotated_anomalies(docs):
    assert {'ITEM-001','ITEM-002','MATH-001','DATE-001','FIELD-001','BASIS-001'}<=rules(docs('B'))

def test_decimal_tolerance(docs):
    d=docs();i=d[1]['data']['items'][0]
    i['price']['value']='0.1';i['quantity']['value']='3';i['amount']['value']='0.30'
    assert 'MATH-001' not in rules(d)
    i['amount']['value']='0.31';assert 'MATH-001' not in rules(d)
    i['amount']['value']='0.311';assert 'MATH-001' in rules(d)
    with pytest.raises(ValueError): audit(d,'NaN')

def test_sum(docs):
    d=docs();d[1]['data']['total']['value']='6001';assert 'MATH-002' in rules(d)

def test_matching_unique_and_ambiguous(docs):
    d=docs();a=d[0]['data']['items'][0];b=d[1]['data']['items'][0]
    b['name']['value']='专业 显示器';b['model']['value']='other';assert matching(a,b)=='可靠'
    b['name']['value']='其他设备';assert matching(a,b)=='待确认'
    assert 'MATCH-001' in rules(d)
    d=docs();d[1]['data']['items'].append(copy.deepcopy(d[1]['data']['items'][0]));assert 'MATCH-001' in rules(d)

def test_unknown_basis_no_direct_price_comparison(docs):
    d=docs();d[2]['data']['tax']['value']=None;d[2]['data']['items'][0]['price']['value']='999'
    r=rules(d);assert 'BASIS-001' in r;assert 'PRICE-001' not in r;assert 'TOTAL-001' not in r

def test_dates(docs):
    d=docs();d[2]['data']['delivery']['value']='2026年11月30日';assert 'DATE-001' not in rules(d)
    d[2]['data']['delivery']['value']='2026-12-01';assert 'DATE-001' in rules(d)
    d[2]['data']['delivery']['value']='签约后30天内';assert 'DATE-001' not in rules(d);assert 'DATE-002' in rules(d)
    assert date_range('2026-02-30') is None

def test_failed_is_not_missing(docs):
    d=docs()
    for x in d: x.update(status='抽取失败',data=None,error='扫描件不支持')
    r=audit(d);assert r['status']=='抽取失败';assert not r['issues'];assert any(c['status']=='抽取失败' for c in r['checks'])
    d.pop();assert 'DOC-001' in rules(d)

def test_buyer_does_not_require_supplier(docs):
    assert docs()[0]['data']['supplier']['value'] is None
    assert 'FIELD-001' not in rules(docs())

def test_evidence_rejects_fabrication(docs):
    d=docs()[0];d['data']['buyer']['quote']='不存在的文本'
    with pytest.raises(ValueError): validate_evidence(d['data'],d['blocks'])

def test_pdf_page_and_docx_table_locations(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer=PdfWriter();page=writer.add_blank_page(width=600,height=800)
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
    stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 20 700 Td (Supplier Quote Model MX27 Quantity 10 Unit 600 Amount 6000) Tj ET')
    page[NameObject('/Contents')]=writer._add_object(stream)
    p=tmp_path/'text.pdf';writer.write(p)
    blocks=parse(p);assert blocks[0]['location']=='PDF 第 1 页';assert 'MX27' in blocks[0]['text']
    from docx import Document
    d=Document();d.add_paragraph('这是一个采购合同文本抽取测试，包含正文和表格。');t=d.add_table(rows=1,cols=2);t.cell(0,0).text='型号';t.cell(0,1).text='MX27';p=tmp_path/'table.docx';d.save(p)
    assert any(b['location']=='DOCX 表格 1 第 1 行' and 'MX27' in b['text'] for b in parse(p))

def test_empty_pdf_is_extraction_failure(tmp_path):
    from pypdf import PdfWriter
    w=PdfWriter();w.add_blank_page(600,800);p=tmp_path/'scan.pdf';w.write(p)
    with pytest.raises(ValueError,match='OCR'):parse(p)

def test_api_lifecycle_and_report(monkeypatch):
    monkeypatch.delenv('MODEL_API_KEY',raising=False)
    c=TestClient(app)
    a=c.post('/api/samples/A').json();b=c.post('/api/samples/B').json()
    c.post(f"/api/tasks/{a['id']}/confirm-fields");c.post(f"/api/tasks/{b['id']}/confirm-fields")
    assert c.post(f"/api/tasks/{a['id']}/audit").json()['result']['status']=='通过'
    b=c.post(f"/api/tasks/{b['id']}/audit").json();iid=b['result']['issues'][0]['id']
    r=c.put(f"/api/tasks/{b['id']}/issues/{iid}",json={'status':'已确认','note':'已比对供应商原文'})
    assert r.json()['result']['issues'][0]['status']=='已确认'
    doc=c.get(f"/api/tasks/{b['id']}/report");assert doc.status_code==200
    from docx import Document
    text='\n'.join(p.text for p in Document(io.BytesIO(doc.content)).paragraphs)
    assert '已比对供应商原文' in text and 'DOCX 段落' in text and '演示' in text
    d=b['documents'][0];payload=d['data'];payload['buyer']['value']='人工修正公司'
    edited=c.put(f"/api/tasks/{b['id']}/documents/{d['id']}",json=payload).json()
    assert edited['result'] is None and len(edited['edits'])==1
    assert c.get(f"/api/tasks/{b['id']}").json()['edits']
    assert c.get('/api/samples/B/download').status_code==200

def test_upload_and_no_key(tmp_path,monkeypatch):
    monkeypatch.delenv('MODEL_API_KEY',raising=False);c=TestClient(app)
    t=c.post('/api/tasks',json={'name':'真实上传测试'}).json();url=f"/api/tasks/{t['id']}"
    assert c.post(url+'/documents',data={'kind':'采购需求'},files={'file':('x.txt',b'abc')}).status_code==400
    assert c.post(url+'/documents',data={'kind':'采购需求'},files={'file':('x.pdf',b'broken')}).json()['documents'][0]['status']=='抽取失败'
    t=c.get(url).json();c.delete(url+'/documents/'+t['documents'][0]['id'])
    p,_=sample('A','采购需求',tmp_path)
    r=c.post(url+'/documents',data={'kind':'采购需求'},files={'file':('real.docx',p.read_bytes())});assert r.status_code==200
    assert c.post(url+'/extract').status_code==409
    for kind in ['供应商报价单','采购合同']:
        path,_=sample('A',kind,tmp_path)
        c.post(url+'/documents',data={'kind':kind},files={'file':(path.name,path.read_bytes())})
    assert c.post(url+'/submit-materials').status_code==200
    r=c.post(url+'/extract').json();d=r['documents'][0]
    assert d['status']=='抽取失败' and d['data'] is None and '密钥' in d['error'] and not d['demo']
    assert c.post(url+'/audit').status_code==409
    a=c.post('/api/samples/A').json();c.post(f"/api/tasks/{a['id']}/confirm-fields")
    assert c.post(f"/api/tasks/{a['id']}/audit?tolerance=-1").status_code==400
